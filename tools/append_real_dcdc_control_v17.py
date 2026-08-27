from pathlib import Path
import shutil

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
DOC_PATH = ROOT / "docs" / "EMS_Modbus双从站读写模拟测试指导书_V1.7.docx"
TEMP_PATH = DOC_PATH.with_suffix(".tmp.docx")

IMAGES = [
    (Path(r"C:\Users\DJ\AppData\Local\Temp\codex-clipboard-9247a628-9442-4d3b-97f1-73feab71709a.png"),
     "图34  启动成功示例：work_state=1、fault_raw=0"),
    (Path(r"C:\Users\DJ\AppData\Local\Temp\codex-clipboard-3dad8a7d-428e-46e0-90db-bccd696a23e8.png"),
     "图35  停机成功示例：work_state=4、fault_raw=0"),
    (Path(r"C:\Users\DJ\AppData\Local\Temp\codex-clipboard-5ccbd984-6fc0-4800-92ea-a94fd41835bc.png"),
     "图36  第三次启动后触发B侧过压：work_state=4、fault_raw=0x0004"),
    (Path(r"C:\Users\DJ\AppData\Local\Temp\codex-clipboard-7a1acfa4-53f7-40ea-82f2-950a6c05cde3.png"),
     "图37  故障后安全停机命令写入与0x03读回均为0x0005"),
    (Path(r"C:\Users\DJ\AppData\Local\Temp\codex-clipboard-8f72c858-4f2a-4aee-9c83-074e24625c87.png"),
     "图38  安全停机命令sequence=9且无新鲜度、故障门禁或总超时拒绝"),
    (Path(r"C:\Users\DJ\AppData\Local\Temp\codex-clipboard-5c71ae4a-a3ba-45e9-af1b-cdd6357f4be9.png"),
     "图39  安全停机后通信继续更新，但B侧过压故障仍锁存"),
    (Path(r"C:\Users\DJ\AppData\Local\Temp\codex-clipboard-8c1af0fe-47aa-4671-b78a-cf73d2403d2e.png"),
     "图40  PC独立读取0x0427～0x0429参数，14次请求无错误"),
]


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text, bold=False, center=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(str(text))
    r.bold = bold
    r.font.size = Pt(9)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    for index, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], header, bold=True, center=True)
        shade_cell(table.rows[0].cells[index], "D9EAF7")
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            set_cell_text(cells[index], value, center=(index == 0))
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Inches(width)
    doc.add_paragraph()
    return table


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(text)


def add_picture(doc, path, caption):
    if not path.exists():
        raise FileNotFoundError(path)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(6.1))
    cp = doc.add_paragraph(caption)
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in cp.runs:
        run.font.size = Pt(9)
    doc.add_paragraph()


def main():
    doc = Document(DOC_PATH)
    marker = "22. 真实DCDC新协议启停与重复性测试"
    if any(p.text.strip() == marker for p in doc.paragraphs):
        print("Section already exists; no change made.")
        return

    doc.add_page_break()
    doc.add_heading(marker, level=1)
    doc.add_paragraph(
        "测试日期：2026年8月11日    设备：真实双向DCDC    从站地址：0xFF    "
        "接口：EMS USART2二线RS485    记录版本：V1.7追加"
    )

    doc.add_heading("22.1 测试目的、边界与安全条件", level=2)
    add_bullet(doc, "验证新协议下0x0402准备、0x0403启停、0x0404工作状态和0x0405故障位的实机表现。")
    add_bullet(doc, "验证DCDC异步写队列、有限重试、总超时、新鲜度门禁、故障启动门禁及写后0x03读回。")
    add_bullet(doc, "原计划连续执行5轮启停；任一轮出现fault_raw非零时立即终止，不再发送启动命令。")
    add_bullet(doc, "本测试仅允许人工监护的Keil Watch触发，不代表TouchGFX或EMS自动控制已经获准接入。")
    add_table(
        doc,
        ["项目", "实测/配置"],
        [
            ["通信参数", "Modbus RTU，9600-8-N-1，DCDC地址0xFF"],
            ["接收方式", "USART2 DMA+IDLE；D-Cache开启；UART DMA地址槽由MPU配置为非缓存"],
            ["写命令保护", "数据新鲜度2 s、最多3次尝试、总时限4 s、写后0x03读回"],
            ["人工触发", "Stop状态准备address/count/values，最后写trigger=0x5AA5，再Run"],
            ["禁止操作", "禁止Run过程中直接修改控制值；故障后禁止再次启动"],
        ],
        [1.35, 5.15],
    )

    doc.add_heading("22.2 新协议与实机状态解释", level=2)
    add_table(
        doc,
        ["地址", "名称", "命令/状态", "当前解释"],
        [
            ["0x0402", "工作模式选择", "写0", "待机/启动准备；写入后实机曾出现work_state=0、fault_raw=0"],
            ["0x0403", "启停控制", "写4", "手动开机，成功后实测work_state=1"],
            ["0x0403", "启停控制", "写5", "安全停机，成功后实测work_state=4、fault_raw=0"],
            ["0x0404", "DCDC工作状态", "读1/4", "1为恒压限流运行；4必须结合0x0405区分停机与故障"],
            ["0x0405", "故障位掩码", "读0/非0", "0表示无故障；0x0004表示B侧过压，作为故障权威依据"],
        ],
        [0.8, 1.2, 1.0, 3.5],
    )
    p = doc.add_paragraph()
    p.add_run("协议差异：").bold = True
    p.add_run(
        "厂家新文档称0x0404只有1和4，但实机在0x0402=0后出现过状态0。状态0的正式语义仍待厂家确认，"
        "当前代码不得仅依据0x0404判断故障，必须同时检查0x0405。"
    )

    doc.add_heading("22.3 启停顺序与前两轮结果", level=2)
    add_table(
        doc,
        ["步骤", "操作", "等待/判定"],
        [
            ["1", "写0x0402=0", "确认写后读回，等待约3 s"],
            ["2", "写0x0403=4", "运行约20 s；要求work_state=1且fault_raw=0"],
            ["3", "写0x0403=5", "等待约10 s；要求work_state=4且fault_raw=0"],
            ["4", "重复上述过程", "计划5轮；故障立即停止"],
        ],
        [0.6, 2.0, 3.9],
    )
    doc.add_paragraph("前两轮均完成准备、启动和停机；下图为成功启停的代表性状态。")
    add_picture(doc, *IMAGES[0])
    add_picture(doc, *IMAGES[1])

    doc.add_heading("22.4 第三次启动B侧过压保护", level=2)
    add_table(
        doc,
        ["观察项", "实测结果", "判定"],
        [
            ["触发阶段", "第三次启动（0x0403=4）后", "重复性测试失败"],
            ["工作状态", "work_state=4", "停止/故障状态"],
            ["故障位", "fault_raw=0x0004", "B侧过压保护"],
            ["保护后B侧电压", "约47.8 V", "是保护后采样，不能代表启动峰值"],
            ["保护后P侧电压", "约0.1 V", "设备已停止功率传输"],
            ["通信状态", "online=1、last_error=0", "故障数据来自持续有效轮询"],
        ],
        [1.25, 2.1, 3.15],
    )
    add_picture(doc, *IMAGES[2])

    doc.add_heading("22.5 故障后的安全停机", level=2)
    add_table(
        doc,
        ["字段", "实测值", "解释"],
        [
            ["address / requested / readback", "0x0403 / 0x0005 / 0x0005", "停机命令写入并读回一致"],
            ["state / function_code", "3 / 0x06", "命令完成；使用单寄存器写"],
            ["attempt_count / readback_confirmed", "1 / 1", "一次成功且完成读回确认"],
            ["last_error / sequence", "0 / 9", "无写错误；第三轮故障后的安全停机为第9条命令"],
            ["三个保护计数", "stale=0、device_fault=0、total_timeout=0", "安全停机未被新鲜度或故障门禁阻断"],
        ],
        [1.65, 1.8, 3.05],
    )
    add_picture(doc, *IMAGES[3])
    add_picture(doc, *IMAGES[4])
    doc.add_paragraph(
        "安全停机不会清除DCDC锁存故障。随后EMS仍在线并持续更新，success_count由0x2E8增至0x2F9，"
        "protocol_error_count保持0x21，但work_state=4、fault_raw=0x0004仍保持。"
    )
    add_picture(doc, *IMAGES[5])

    doc.add_heading("22.6 输出设定与保护阈值核对", level=2)
    doc.add_paragraph(
        "为避免双主站冲突，EMS停止/退出主站访问后，由PC Modbus Poll经USB-RS485作为唯一主站读取"
        "0x0427～0x0429，扫描周期500 ms；截图显示Tx=14、Err=0。"
    )
    add_table(
        doc,
        ["地址", "原始值", "工程量/含义"],
        [
            ["0x0427", "540", "B侧输出电压设定54.0 V"],
            ["0x0428", "510", "P侧输出电压设定51.0 V"],
            ["0x0429", "580", "B侧过压保护阈值58.0 V"],
        ],
        [1.0, 1.25, 4.25],
    )
    add_picture(doc, *IMAGES[6])

    doc.add_heading("22.7 验收结论与后续限制", level=2)
    add_table(
        doc,
        ["验收项", "结果", "结论"],
        [
            ["基本启停与读回", "通过", "写队列、0x06及0x03读回基本可用"],
            ["故障启动门禁/安全停机", "通过", "停机可在故障状态执行；故障锁存不会被误清除"],
            ["5轮启停重复性", "未通过", "第三次启动触发B侧过压，第四、第五轮中止"],
            ["TouchGFX/EMS自动控制放行", "不通过", "保护原因和偶发通信问题均未闭环"],
        ],
        [2.0, 1.0, 3.5],
    )
    callout = doc.add_table(rows=1, cols=1)
    callout.style = "Table Grid"
    set_cell_text(
        callout.cell(0, 0),
        "最终结论：第三次启动触发真实DCDC B侧过压保护，本轮5次启停测试判定未通过并中止。"
        "在厂家确认54.0 V输出设定、58.0 V保护阈值、状态0语义及当前接线/负载适配性，并使用示波器或"
        "具有最大值保持能力的仪表检查启动瞬态之前，禁止再次启动、禁止提高保护阈值、禁止接入TouchGFX或EMS自动控制。",
        bold=True,
    )
    shade_cell(callout.cell(0, 0), "FCE4D6")
    doc.add_paragraph()
    add_bullet(doc, "保持DCDC停机并保留当前截图、参数和sequence=9记录。")
    add_bullet(doc, "向厂家确认0x0427=540、0x0429=580是否适合当前工况，以及work_state=0的正式含义。")
    add_bullet(doc, "使用示波器或最大值保持仪表捕获启动瞬间B侧电压；500 ms Modbus轮询不能排除短时过冲。")
    add_bullet(doc, "原因修正后从第1轮重新执行至少5轮完整启停，不从第4轮续测。")

    doc.core_properties.modified = doc.core_properties.modified
    doc.save(TEMP_PATH)
    shutil.move(str(TEMP_PATH), str(DOC_PATH))
    print(f"Updated {DOC_PATH}")


if __name__ == "__main__":
    main()
