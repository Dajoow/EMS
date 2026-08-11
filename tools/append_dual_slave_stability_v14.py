from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(r"C:\Users\DJ\Desktop\Hdu_EMS_H750")
SOURCE = ROOT / "docs" / "EMS_Modbus双从站读写模拟测试指导书_V1.3.docx"
OUTPUT = ROOT / "docs" / "EMS_Modbus双从站读写模拟测试指导书_V1.4.docx"
IMG_DCDC = Path(r"C:\Users\DJ\AppData\Local\Temp\codex-clipboard-e08fe7df-00c2-4fa0-8d70-d9aecdc58f8a.png")
IMG_ACDC = Path(r"C:\Users\DJ\AppData\Local\Temp\codex-clipboard-a603ba43-222d-425d-9ab6-6b45bffd0cf6.png")
IMG_SLAVE = Path(r"C:\Users\DJ\AppData\Local\Temp\codex-clipboard-ad169c9c-a568-4742-94a8-1b824d3f2b1d.png")


def set_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_margins(cell, top=100, start=120, bottom=100, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def format_table(table, widths_cm):
    table.autofit = False
    for row_index, row in enumerate(table.rows):
        for col_index, cell in enumerate(row.cells):
            cell.width = Cm(widths_cm[col_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_margins(cell)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.05
                for run in paragraph.runs:
                    run.font.name = "Microsoft YaHei"
                    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "微软雅黑")
                    run.font.size = Pt(9)
            if row_index == 0:
                set_shading(cell, "D9EAF7")
                for run in cell.paragraphs[0].runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor(31, 78, 121)


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.add_run(text)


def add_picture(doc, path, width_cm, caption):
    if not path.exists():
        return
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run().add_picture(str(path), width=Cm(width_cm))
    paragraph = doc.add_paragraph(caption)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(8)


doc = Document(SOURCE)
doc.add_page_break()
doc.add_heading("18. ACDC与DCDC双从站正常通信稳定性测试", level=1)

p = doc.add_paragraph()
p.add_run("测试日期：").bold = True
p.add_run("2026年7月21日    ")
p.add_run("测试方式：").bold = True
p.add_run("PC Modbus Slave模拟双从站    ")
p.add_run("记录版本：").bold = True
p.add_run("V1.4")

doc.add_heading("18.1 测试目标", level=2)
add_bullet(doc, "验证EMS通过同一条二线RS485总线，作为唯一Modbus RTU主站同时轮询ACDC和DCDC。")
add_bullet(doc, "验证ACDC从站地址33（0x21）与DCDC从站地址1（0x01）能够正确区分，数据不会串站。")
add_bullet(doc, "验证两个从站在长时间连续轮询过程中保持在线，且不产生超时和协议错误。")
add_bullet(doc, "验证分别修改ACDC和DCDC模拟寄存器后，EMS工程量能够独立、正确更新。")

doc.add_heading("18.2 从站及寄存器窗口配置", level=2)
table = doc.add_table(rows=1, cols=6)
table.style = "Table Grid"
for index, value in enumerate(("设备", "Slave ID", "功能码", "起始地址(Dec)", "数量", "用途")):
    table.rows[0].cells[index].text = value
rows = [
    ("ACDC", "33", "03", "4096", "13", "母排、交流电压电流及频率"),
    ("ACDC", "33", "03", "4130", "2", "32位直流负载功率"),
    ("ACDC", "33", "03", "4352", "4", "整流模块汇总"),
    ("ACDC", "33", "03", "20482", "12", "基础交流告警"),
    ("DCDC", "1", "03", "1028", "11", "状态、故障及温度"),
    ("DCDC", "1", "03", "1051", "6", "功率、电压及电流"),
    ("DCDC", "1", "03", "1026", "2", "控制区，仅写入后读回"),
    ("DCDC", "1", "03", "1063", "14", "参数区，仅写入后读回"),
]
for values in rows:
    cells = table.add_row().cells
    for index, value in enumerate(values):
        cells[index].text = value
format_table(table, [2.2, 2.0, 1.8, 3.0, 1.6, 5.6])

p = doc.add_paragraph()
p.add_run("配置说明：").bold = True
p.add_run(
    "全部窗口使用03 Holding Register和十进制原始地址；PLC Addresses (Base 1)、Read/Write Disabled、"
    "Skip response、CRC/LRC error和Return exception均未勾选。DCDC地址1026和1063窗口在未执行写命令时Rx=0属于正常现象。"
)

doc.add_heading("18.3 动态工程量验证", level=2)
table = doc.add_table(rows=1, cols=5)
table.style = "Table Grid"
for index, value in enumerate(("设备", "模拟寄存器", "填写值", "EMS字段", "实测结果")):
    table.rows[0].cells[index].text = value
rows = [
    ("ACDC", "4096 / 0x1000", "500", "dc_bus_voltage_v", "50.0V，正确"),
    ("DCDC", "1053 / 0x041D", "500", "b_voltage_v", "50.0V，正确"),
]
for values in rows:
    cells = table.add_row().cells
    for index, value in enumerate(values):
        cells[index].text = value
format_table(table, [2.2, 3.5, 2.0, 4.8, 3.7])

p = doc.add_paragraph()
p.add_run("判定：").bold = True
p.add_run("ACDC和DCDC模拟值分别只更新各自数据结构，未观察到地址混淆或跨从站数据串扰。")

doc.add_heading("18.4 长时间运行结果", level=2)
table = doc.add_table(rows=1, cols=4)
table.style = "Table Grid"
for index, value in enumerate(("检查项", "ACDC实测", "DCDC实测", "结论")):
    table.rows[0].cells[index].text = value
rows = [
    ("online", "1", "1", "均在线"),
    ("consecutive_failures", "0", "0", "无连续失败"),
    ("last_error", "0", "0", "当前无错误"),
    ("success_count", "0x1635 = 5685", "0x1635 = 5685", "双方完成5685轮完整轮询"),
    ("timeout_count", "0", "0", "无响应超时"),
    ("protocol_error_count", "0", "0", "无协议错误"),
    ("last_update_tick", "0x00504AB0", "0x00504B3F", "双方持续更新"),
]
for values in rows:
    cells = table.add_row().cells
    for index, value in enumerate(values):
        cells[index].text = value
format_table(table, [4.1, 4.1, 4.1, 4.1])

p = doc.add_paragraph()
p.add_run("Rx计数说明：").bold = True
p.add_run(
    "DCDC两个周期读取窗口均为5685次，与DCDC success_count一致。ACDC四个窗口显示约5892～5948次累计请求，"
    "其差值包含窗口创建时间不同及前期单独调试产生的历史计数；ACDC success_count=5685表明本轮已经连续完成5685次四数据块全部成功的完整轮询。"
)

add_picture(doc, IMG_DCDC, 10.5, "图12  DCDC连续5685轮后保持在线，超时和协议错误均为0")
add_picture(doc, IMG_ACDC, 10.5, "图13  ACDC连续5685轮后保持在线，工程量正确且错误计数为0")
doc.add_page_break()
add_picture(doc, IMG_SLAVE, 16.2, "图14  Modbus Slave同时配置ACDC地址33与DCDC地址1的全部测试窗口")

doc.add_heading("18.5 验收结论", level=2)
conclusion = doc.add_table(rows=1, cols=1)
conclusion.style = "Table Grid"
cell = conclusion.cell(0, 0)
cell.text = (
    "最终结论：EMS通过同一个二线RS485接口，作为Modbus RTU主站同时访问ACDC从站33和DCDC从站1。"
    "两个从站均完成5685轮完整轮询，online=1、last_error=0、consecutive_failures=0、timeout_count=0、"
    "protocol_error_count=0；动态工程量分别正确更新且没有数据串扰。因此，ACDC+DCDC双从站正常通信及稳定性测试通过。"
)
set_shading(cell, "E2F0D9")
set_margins(cell, top=170, start=190, bottom=170, end=190)
for run in cell.paragraphs[0].runs:
    run.bold = True
    run.font.name = "Microsoft YaHei"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(55, 86, 35)

doc.add_heading("18.6 阶段边界与后续事项", level=2)
add_bullet(doc, "本结论适用于PC Modbus Slave模拟环境；真实ACDC和真实DCDC的现场只读联调尚未执行。")
add_bullet(doc, "DCDC 0x06和0x10写入及写后读回已在前述章节通过，但真实设备启停控制仍应在安全互锁完成后进行。")
add_bullet(doc, "此前Skip response期间出现的电池页面异常问题仍维持“暂时待定”，本轮正常通信稳定性测试未再出现通信错误。")
add_bullet(doc, "下一阶段建议先连接真实设备进行只读核对，再分别验证单从站断线不会阻塞另一从站。")

doc.save(OUTPUT)
print(OUTPUT)
