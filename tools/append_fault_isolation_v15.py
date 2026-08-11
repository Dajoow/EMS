from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(r"C:\Users\DJ\Desktop\Hdu_EMS_H750")
SOURCE = ROOT / "docs" / "EMS_Modbus双从站读写模拟测试指导书_V1.4.docx"
OUTPUT = ROOT / "docs" / "EMS_Modbus双从站读写模拟测试指导书_V1.5.docx"

IMAGES = {
    "acdc_offline": Path(r"C:\Users\DJ\AppData\Local\Temp\codex-clipboard-415a942b-cd2a-4745-b8fd-1250454ea2a1.png"),
    "dcdc_during_acdc_fault": Path(r"C:\Users\DJ\AppData\Local\Temp\codex-clipboard-8c162820-cd87-4ff3-a9c8-ddf6910f5fc9.png"),
    "acdc_fault_windows": Path(r"C:\Users\DJ\AppData\Local\Temp\codex-clipboard-1f58aacd-baaf-42df-a90e-39bb62d4a4d1.png"),
    "acdc_recovered": Path(r"C:\Users\DJ\AppData\Local\Temp\codex-clipboard-df99398c-821b-4b8b-9ac6-ce4b8e089d86.png"),
    "acdc_recovery_windows": Path(r"C:\Users\DJ\AppData\Local\Temp\codex-clipboard-7a1a8c51-3454-4da1-8e05-3edb49670d42.png"),
    "dcdc_offline": Path(r"C:\Users\DJ\AppData\Local\Temp\codex-clipboard-343750d1-2463-4a29-85b4-f7f044adbb35.png"),
    "acdc_during_dcdc_fault": Path(r"C:\Users\DJ\AppData\Local\Temp\codex-clipboard-0f8d3024-0a80-4ed6-98f6-c03f06635708.png"),
    "dcdc_fault_windows": Path(r"C:\Users\DJ\AppData\Local\Temp\codex-clipboard-fec62fa0-4594-464c-9aa6-7602d6d7697d.png"),
    "both_recovered": Path(r"C:\Users\DJ\AppData\Local\Temp\codex-clipboard-5bb6e507-1f07-451b-b911-84279254ad0d.png"),
    "dcdc_recovery_windows": Path(r"C:\Users\DJ\AppData\Local\Temp\codex-clipboard-51ee6477-09f7-41bd-a13a-0a41fe4cf993.png"),
}


def set_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_margins(cell, top=100, start=120, bottom=100, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def format_table(table, widths_cm):
    table.autofit = False
    for row_index, row in enumerate(table.rows):
        for col_index, cell in enumerate(row.cells):
            cell.width = Cm(widths_cm[col_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_margins(cell)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.05
                for run in paragraph.runs:
                    run.font.name = "Microsoft YaHei"
                    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "微软雅黑")
                    run.font.size = Pt(9)
            if row_index == 0:
                set_shading(cell, "D9EAF7")
                for run in cell.paragraphs[0].runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor(31, 78, 121)


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.add_run(text)


def add_picture(doc, key, width_cm, caption):
    path = IMAGES[key]
    if not path.exists():
        raise FileNotFoundError(path)
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run().add_picture(str(path), width=Cm(width_cm))
    paragraph = doc.add_paragraph(caption)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(8)


def add_result_table(doc, rows):
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for index, value in enumerate(("观察项", "故障前/历史基线", "故障或恢复实测", "判定")):
        table.rows[0].cells[index].text = value
    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = value
    format_table(table, [4.0, 4.1, 4.6, 3.7])


doc = Document(SOURCE)
doc.core_properties.title = "EMS Modbus RTU 双从站 PC 模拟测试指导书 V1.5"
doc.add_page_break()
doc.add_heading("19. ACDC/DCDC异常隔离与自动恢复测试", level=1)

p = doc.add_paragraph()
p.add_run("测试日期：").bold = True
p.add_run("2026年7月22日    ")
p.add_run("测试方式：").bold = True
p.add_run("PC Modbus Slave修改从站地址，模拟持续失联    ")
p.add_run("记录版本：").bold = True
p.add_run("V1.5")

doc.add_heading("19.1 测试目的与方法", level=2)
add_bullet(doc, "验证ACDC持续失联时，EMS能够判定ACDC离线，同时DCDC保持在线并继续更新。")
add_bullet(doc, "验证DCDC持续失联时，EMS能够判定DCDC离线，同时ACDC保持在线且错误计数不增加。")
add_bullet(doc, "验证故障窗口恢复原从站地址后，目标从站无需复位EMS即可自动恢复通信。")
add_bullet(doc, "故障注入期间保持EMS运行和串口连接；ACDC四个窗口由ID 33改为ID 30，DCDC四个窗口由ID 1改为ID 2。")

p = doc.add_paragraph()
p.add_run("工具行为说明：").bold = True
p.add_run(
    "Modbus Slave的Skip response为间歇丢包模拟，本次实测约每10次请求跳过1次响应，不能形成连续3次失败。"
    "因此持续断线测试改用临时修改全部目标从站窗口ID的方法。窗口重新定义会使该窗口Rx从0重新计数。"
)

doc.add_heading("19.2 ACDC持续断线与DCDC隔离", level=2)
p = doc.add_paragraph()
p.add_run("操作：").bold = True
p.add_run("保持通信运行，将ACDC四个周期窗口从ID 33依次改为ID 30；全部修改完成后观察通信状态。")

add_result_table(doc, [
    ("ACDC online", "1", "0", "通过：离线"),
    ("ACDC consecutive_failures", "0", "0x30 = 48", "超过离线门限3"),
    ("ACDC last_error", "0", "0xFFFE = -2", "持续无响应超时"),
    ("ACDC success_count", "0x432 = 1074", "保持0x432", "故障期间停止成功累计"),
    ("ACDC timeout_count", "历史累计", "0x89 = 137并持续增加", "超时检测有效"),
    ("DCDC online/last_error", "1 / 0", "1 / 0", "ACDC故障未影响DCDC"),
    ("DCDC success_count", "持续增加", "0x514 = 1300", "DCDC继续轮询"),
    ("DCDC错误计数", "0", "timeout=0，protocol=0", "无串扰"),
])

add_picture(doc, "acdc_offline", 10.8, "图15  ACDC持续失联后online=0、last_error=-2")
add_picture(doc, "dcdc_during_acdc_fault", 10.8, "图16  ACDC失联期间DCDC保持在线且错误计数为0")
doc.add_page_break()
add_picture(doc, "acdc_fault_windows", 16.2, "图17  ACDC窗口改为ID 30，DCDC地址1继续接收请求")

doc.add_heading("19.3 ACDC自动恢复", level=2)
p = doc.add_paragraph()
p.add_run("操作：").bold = True
p.add_run("保持EMS运行，将ACDC四个窗口依次恢复为ID 33；全部恢复后开始判定。")

add_result_table(doc, [
    ("ACDC online", "断线时0", "1", "恢复在线"),
    ("ACDC consecutive_failures", "0x30", "0", "当前失败清除"),
    ("ACDC last_error", "-2", "0", "当前错误清除"),
    ("ACDC success_count", "0x432 = 1074", "0x44D = 1101", "增加27轮"),
    ("ACDC last_update_tick", "断线时冻结", "重新变化", "数据恢复更新"),
    ("DCDC状态", "在线", "始终在线且错误计数为0", "隔离有效"),
])

p = doc.add_paragraph()
p.add_run("Rx差异说明：").bold = True
p.add_run(
    "ACDC四个窗口逐个恢复，靠前的数据块更早开始计数，因此恢复截图中的累计Rx不同；全部恢复后只要求各窗口增量持续增加，"
    "不要求历史累计值相等。恢复过程中产生的少量超时/协议错误属于窗口ID尚未全部恢复时的过渡记录。"
)

add_picture(doc, "acdc_recovered", 10.8, "图18  ACDC无需复位恢复在线并重新累计成功轮次")
add_picture(doc, "acdc_recovery_windows", 16.2, "图19  ACDC四个窗口恢复ID 33后的通信状态")

doc.add_heading("19.4 DCDC持续断线与ACDC隔离", level=2)
p = doc.add_paragraph()
p.add_run("操作：").bold = True
p.add_run(
    "确认所有窗口的Skip response、CRC/LRC error和Return exception均取消后，将DCDC四个窗口由ID 1改为ID 2。"
    "此前曾因ACDC1残留Skip response导致ACDC超时计数增加；清除该选项后重新执行了干净复测。"
)

add_result_table(doc, [
    ("DCDC online", "1", "0", "通过：离线"),
    ("DCDC consecutive_failures", "0", "0x19 = 25", "超过离线门限3"),
    ("DCDC last_error", "0", "0xFFFE = -2", "持续无响应超时"),
    ("DCDC success_count", "0x749 = 1865", "保持0x749", "故障期间停止成功累计"),
    ("DCDC timeout_count", "历史累计", "0x25 = 37", "超时路径有效"),
    ("ACDC online/last_error", "1 / 0", "1 / 0", "DCDC故障未影响ACDC"),
    ("ACDC success_count", "0x669 = 1641", "持续增加", "ACDC继续轮询"),
    ("ACDC错误计数", "timeout=0xBC，protocol=0x72", "保持不变", "干净复测无串扰"),
])

add_picture(doc, "dcdc_offline", 10.8, "图20  DCDC持续失联后online=0、last_error=-2")
add_picture(doc, "acdc_during_dcdc_fault", 10.8, "图21  DCDC失联期间ACDC保持在线且错误计数不增加")
doc.add_page_break()
add_picture(doc, "dcdc_fault_windows", 16.2, "图22  DCDC窗口改为ID 2，ACDC地址33继续接收请求")

doc.add_heading("19.5 DCDC自动恢复", level=2)
p = doc.add_paragraph()
p.add_run("操作：").bold = True
p.add_run("保持EMS运行，将DCDC四个窗口依次恢复为ID 1；全部恢复后观察两个数据结构。")

add_result_table(doc, [
    ("DCDC online", "断线时0", "1", "恢复在线"),
    ("DCDC consecutive_failures", "0x19", "0", "当前失败清除"),
    ("DCDC last_error", "-2", "0", "当前错误清除"),
    ("DCDC success_count", "0x749 = 1865", "0x758 = 1880", "增加15轮"),
    ("DCDC last_update_tick", "断线时冻结", "重新变化", "数据恢复更新"),
    ("ACDC success_count", "0x669 = 1641", "0x691 = 1681", "继续增加40轮"),
    ("ACDC错误计数", "timeout=0xBC，protocol=0x72", "保持不变", "反向隔离通过"),
])

p = doc.add_paragraph()
p.add_run("历史计数说明：").bold = True
p.add_run(
    "DCDC恢复截图中的timeout_count与protocol_error_count包含逐个恢复窗口ID时产生的过渡错误；"
    "全部窗口恢复后online=1、last_error=0、consecutive_failures=0且success_count持续增加，表明自动恢复成功。"
)

add_picture(doc, "both_recovered", 13.5, "图23  DCDC恢复后两个从站均在线，ACDC错误计数保持不变")
add_picture(doc, "dcdc_recovery_windows", 16.2, "图24  DCDC窗口恢复ID 1后的通信状态")

doc.add_heading("19.6 验收结论", level=2)
table = doc.add_table(rows=1, cols=3)
table.style = "Table Grid"
for index, value in enumerate(("验收项", "结果", "结论")):
    table.rows[0].cells[index].text = value
rows = [
    ("ACDC持续失联检测", "online=0，last_error=-2", "通过"),
    ("ACDC故障时DCDC隔离", "DCDC持续在线，错误计数为0", "通过"),
    ("ACDC自动恢复", "无需复位，成功计数重新增加", "通过"),
    ("DCDC持续失联检测", "online=0，last_error=-2", "通过"),
    ("DCDC故障时ACDC隔离", "ACDC持续在线，干净复测错误计数不变", "通过"),
    ("DCDC自动恢复", "无需复位，成功计数重新增加", "通过"),
]
for values in rows:
    cells = table.add_row().cells
    for index, value in enumerate(values):
        cells[index].text = value
format_table(table, [7.0, 6.2, 3.2])

conclusion = doc.add_table(rows=1, cols=1)
conclusion.style = "Table Grid"
cell = conclusion.cell(0, 0)
cell.text = (
    "最终结论：EMS通过同一条二线RS485总线轮询ACDC和DCDC时，任一从站持续失联均只影响该从站；"
    "另一从站保持在线并继续更新。故障从站连续失败达到门限后正确离线，恢复原地址后无需复位EMS即可自动恢复。"
    "ACDC→DCDC和DCDC→ACDC两个方向的持续断线隔离与自动恢复测试均通过。"
)
set_shading(cell, "E2F0D9")
set_margins(cell, top=170, start=190, bottom=170, end=190)
for run in cell.paragraphs[0].runs:
    run.bold = True
    run.font.name = "Microsoft YaHei"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(55, 86, 35)

doc.add_heading("19.7 后续测试", level=2)
add_bullet(doc, "分别对ACDC和DCDC执行CRC/LRC错误注入，验证protocol_error_count、错误恢复和另一从站隔离。")
add_bullet(doc, "异常注入与恢复至少重复5轮，记录每轮离线、恢复时间、错误计数和是否复现电池页面异常。")
add_bullet(doc, "继续监视Client_Sd[0].BAT_VOL[]、BAT_SOC[]、BAT_TMP[]，定位此前DCDC Skip response期间电池页面异常值问题。")
add_bullet(doc, "真实设备接入前关闭DCDC_MODBUS_PC_TEST_ENABLE，并先执行ACDC/DCDC只读联调。")

doc.save(OUTPUT)
print(OUTPUT)
