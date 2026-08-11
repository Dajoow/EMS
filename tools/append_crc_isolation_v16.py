from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(r"C:\Users\DJ\Desktop\Hdu_EMS_H750")
SOURCE = ROOT / "docs" / "EMS_Modbus双从站读写模拟测试指导书_V1.5.docx"
OUTPUT = ROOT / "docs" / "EMS_Modbus双从站读写模拟测试指导书_V1.6.docx"

IMAGES = {
    "acdc_status": Path(r"C:\Users\DJ\AppData\Local\Temp\codex-clipboard-b981fdc8-274c-4e4a-9fa4-fc183a6b226c.png"),
    "acdc_windows": Path(r"C:\Users\DJ\AppData\Local\Temp\codex-clipboard-331b3de2-5dd3-4c64-b596-0d7a4dbaa8fb.png"),
    "dcdc_status": Path(r"C:\Users\DJ\AppData\Local\Temp\codex-clipboard-611e9093-7f7f-4c29-afe0-8aaabebcd77f.png"),
    "dcdc_windows": Path(r"C:\Users\DJ\AppData\Local\Temp\codex-clipboard-21952eb5-5bee-485d-b163-39e8c50b91a5.png"),
}


def set_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_margins(cell, top=100, start=120, bottom=100, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def format_table(table, widths_cm):
    table.autofit = False
    for row_index, row in enumerate(table.rows):
        for col_index, cell in enumerate(row.cells):
            cell.width = Cm(widths_cm[col_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_margins(cell)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.05
                for run in paragraph.runs:
                    run.font.name = "Microsoft YaHei"
                    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "微软雅黑")
                    run.font.size = Pt(9)
            if row_index == 0:
                set_shading(cell, "D9EAF7")
                for run in cell.paragraphs[0].runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor(31, 78, 121)


def add_bullet(doc, text):
    doc.add_paragraph(text, style="List Bullet")


def add_picture(doc, key, width_cm, caption):
    path = IMAGES[key]
    if not path.exists():
        raise FileNotFoundError(path)
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run().add_picture(str(path), width=Cm(width_cm))
    paragraph = doc.add_paragraph(caption)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(8)


def add_result_table(doc, rows):
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for index, value in enumerate(("观察项", "注入前/历史值", "注入及恢复结果", "判定")):
        table.rows[0].cells[index].text = value
    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = value
    format_table(table, [4.0, 4.1, 4.8, 3.5])


doc = Document(SOURCE)
doc.core_properties.title = "EMS Modbus RTU 双从站 PC 模拟测试指导书 V1.6"
doc.add_page_break()
doc.add_heading("20. ACDC/DCDC CRC异常隔离与自动恢复测试", level=1)

p = doc.add_paragraph()
p.add_run("测试日期：").bold = True
p.add_run("2026年7月28日    ")
p.add_run("测试工具：").bold = True
p.add_run("Witte Software Modbus Slave + USB转RS485    ")
p.add_run("记录版本：").bold = True
p.add_run("V1.6")

doc.add_heading("20.1 测试目的与判定规则", level=2)
add_bullet(doc, "验证EMS能够把CRC错误识别为协议/帧错误，而不是响应超时。")
add_bullet(doc, "验证ACDC发生CRC错误时DCDC仍正常，DCDC发生CRC错误时ACDC仍正常。")
add_bullet(doc, "验证取消CRC错误注入后，目标从站无需复位EMS即可自动恢复稳定通信。")
add_bullet(doc, "protocol_error_count和timeout_count为历史累计值；恢复成功不要求历史值清零，只要求错误停止增长且当前状态恢复正常。")

p = doc.add_paragraph()
p.add_run("注入方法：").bold = True
p.add_run(
    "保持EMS和Modbus Slave在线运行，在目标周期读取窗口勾选Insert CRC/LRC error，持续60秒；"
    "随后取消勾选并继续运行60秒。ACDC使用ID 33、地址4096窗口，DCDC使用ID 1、地址1028窗口。"
    "Skip response和Return exception保持未勾选。"
)

doc.add_heading("20.2 ACDC CRC错误注入、隔离与恢复", level=2)
add_result_table(doc, [
    ("ACDC protocol_error_count", "0x0C = 12", "注入后0x11 = 17；取消注入后停止增加", "识别5次CRC错误"),
    ("ACDC timeout_count", "0", "保持0", "未误判为超时"),
    ("ACDC当前状态", "online=1", "online=1，failures=0，last_error=0", "自动恢复正常"),
    ("ACDC success_count", "0xC4 = 196", "恢复观察时0x11D = 285并继续增加", "有效帧继续处理"),
    ("DCDC当前状态", "online=1", "online=1，failures=0，last_error=0", "未受ACDC错误影响"),
    ("DCDC错误计数", "timeout=0，protocol=0", "全程保持0", "正向隔离通过"),
])

add_picture(doc, "acdc_status", 11.0, "图25  ACDC CRC错误计数增加，DCDC全程保持正常")
add_picture(doc, "acdc_windows", 16.2, "图26  ACDC CRC测试期间两个从站窗口保持轮询")

doc.add_heading("20.3 DCDC CRC错误注入、隔离与恢复", level=2)
add_result_table(doc, [
    ("DCDC protocol_error_count", "0", "注入60秒后0x07 = 7；取消注入60秒后保持0x07", "CRC检测与停止增长均正常"),
    ("DCDC timeout_count", "0", "保持0", "未误判为超时"),
    ("DCDC当前状态", "online=1", "online=1，failures=0，last_error=0", "自动恢复正常"),
    ("DCDC success_count", "注入前持续增加", "恢复观察时0x180 = 384并继续增加", "有效帧继续处理"),
    ("ACDC当前状态", "online=1", "online=1，failures=0，last_error=0", "未受DCDC错误影响"),
    ("ACDC timeout_count", "0", "保持0", "无超时串扰"),
    ("ACDC protocol_error_count", "历史值0x11 = 17", "全程保持0x11", "反向隔离通过"),
    ("ACDC success_count", "持续增加", "恢复观察时0x1A6 = 422", "ACDC持续轮询"),
])

add_picture(doc, "dcdc_status", 11.0, "图27  DCDC CRC计数保持0x07，ACDC错误计数未变化")
add_picture(doc, "dcdc_windows", 16.2, "图28  取消DCDC CRC注入后双从站继续正常轮询")

doc.add_heading("20.4 验收结论", level=2)
table = doc.add_table(rows=1, cols=3)
table.style = "Table Grid"
for index, value in enumerate(("验收项", "实测结果", "结论")):
    table.rows[0].cells[index].text = value
for values in [
    ("ACDC CRC错误识别", "protocol增加5，timeout保持0", "通过"),
    ("ACDC错误时DCDC隔离", "DCDC在线且错误计数为0", "通过"),
    ("ACDC取消注入后恢复", "协议错误停止增加，成功计数继续", "通过"),
    ("DCDC CRC错误识别", "protocol增加7，timeout保持0", "通过"),
    ("DCDC错误时ACDC隔离", "ACDC在线，历史错误计数不变", "通过"),
    ("DCDC取消注入后恢复", "protocol保持0x07，成功计数继续", "通过"),
]:
    cells = table.add_row().cells
    for index, value in enumerate(values):
        cells[index].text = value
format_table(table, [7.0, 6.2, 3.2])

conclusion = doc.add_table(rows=1, cols=1)
conclusion.style = "Table Grid"
cell = conclusion.cell(0, 0)
cell.text = (
    "最终结论：ACDC和DCDC双向CRC异常隔离与自动恢复测试通过。"
    "EMS能够把损坏的响应帧计入protocol_error_count，未错误计入timeout_count；"
    "单个从站的CRC异常不会阻塞或污染另一从站，取消注入后无需复位EMS即可恢复稳定通信。"
)
set_shading(cell, "E2F0D9")
set_margins(cell, top=170, start=190, bottom=170, end=190)
for run in cell.paragraphs[0].runs:
    run.bold = True
    run.font.name = "Microsoft YaHei"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(55, 86, 35)

doc.add_heading("20.5 下一项测试", level=2)
add_bullet(doc, "分别对ACDC和DCDC返回Modbus异常响应，使用软件界面提供的Exception 06（Slave Device Busy）。")
add_bullet(doc, "确认目标从站protocol_error_count增加、timeout_count不增加、另一从站保持正常。")
add_bullet(doc, "取消Return exception后继续运行60秒，确认目标从站无需复位即可恢复。")

doc.save(OUTPUT)
print(OUTPUT)
