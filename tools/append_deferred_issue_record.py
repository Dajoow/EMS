from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(r"C:\Users\DJ\Desktop\Hdu_EMS_H750")
SOURCE = ROOT / "docs" / "EMS_Modbus双从站读写模拟测试指导书_V1.2.docx"
OUTPUT = ROOT / "docs" / "EMS_Modbus双从站读写模拟测试指导书_V1.3.docx"
PHOTO = Path(r"C:\Users\DJ\AppData\Local\Temp\codex-clipboard-83bb4c46-2ba3-4bf8-b74a-356383a320ae.png")


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def margins(cell, top=100, start=120, bottom=100, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def format_table(table, widths):
    table.autofit = False
    for row_no, row in enumerate(table.rows):
        for col_no, cell in enumerate(row.cells):
            cell.width = Cm(widths[col_no])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            margins(cell)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.name = "Microsoft YaHei"
                    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "微软雅黑")
                    run.font.size = Pt(9)
            if row_no == 0:
                shade(cell, "D9EAF7")
                for run in cell.paragraphs[0].runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor(31, 78, 121)


doc = Document(SOURCE)
doc.add_heading("17.7 资源余量复核与问题挂起记录", level=2)

p = doc.add_paragraph()
p.add_run("复核日期：").bold = True
p.add_run("2026年7月21日。")
p.add_run("处理决定：").bold = True
p.add_run("本问题暂时待定，不在当前阶段修改工程代码；保留现场和诊断数据，后续结合Client_Sd[]底层数组观察继续定位。")

table = doc.add_table(rows=1, cols=4)
table.style = "Table Grid"
for index, value in enumerate(("监测项目", "实测值", "解释", "结论")):
    table.rows[0].cells[index].text = value

rows = [
    ("FreeHeapSize", "263216", "当前空闲堆空间充足。", "正常"),
    ("MiniFreeHeapSize", "261488", "历史最低空闲堆仍保持较高水平。", "未发现堆耗尽"),
    ("modbus_Process_minStack", "399", "任务初始栈深度512，历史最大使用量约113，剩余约78%。", "基本排除Modbus任务栈溢出"),
    ("CAN_Rev_Thread_minStack", "99", "余量明显小于Modbus任务，但当前仍非零。", "后续持续观察"),
    ("TCP_Rev_Task_minStack", "183", "任务仍有可用栈余量。", "正常"),
    ("GFX_Task_minStack", "1101", "TouchGFX任务栈余量充足。", "正常"),
]

for values in rows:
    cells = table.add_row().cells
    for index, value in enumerate(values):
        cells[index].text = value
format_table(table, [4.0, 2.4, 7.0, 3.0])

if PHOTO.exists():
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(PHOTO), width=Cm(8.3))
    p = doc.add_paragraph("图11  FreeRTOS堆空间及各任务历史最小剩余栈实测")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

p = doc.add_paragraph()
r = p.add_run("阶段判断：")
r.bold = True
r.font.color.rgb = RGBColor(31, 78, 121)
p.add_run(
    "现有数据不支持“Modbus任务栈溢出导致界面异常”的判断。"
    "Skip response期间电池页面异常值的根因仍未确认，后续排查重点保留为："
    "Client_Sd[]是否被改写、CAN接收长度和数组边界、CAN写入与TouchGFX读取的并发一致性。"
)

callout = doc.add_table(rows=1, cols=1)
callout.style = "Table Grid"
cell = callout.cell(0, 0)
cell.text = (
    "问题状态：暂时待定（Deferred）。当前不阻塞PC端Modbus正常读写框架测试记录，"
    "但在连接真实DCDC并开放启停控制之前，必须重新执行异常隔离和恢复测试。"
)
shade(cell, "FFF2CC")
margins(cell, top=150, start=180, bottom=150, end=180)
for run in cell.paragraphs[0].runs:
    run.bold = True
    run.font.name = "Microsoft YaHei"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(127, 96, 0)

doc.save(OUTPUT)
print(OUTPUT)
