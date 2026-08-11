from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(r"C:\Users\DJ\Desktop\Hdu_EMS_H750")
SOURCE = ROOT / "docs" / "EMS_Modbus双从站读写模拟测试指导书_V1.6.docx"
OUTPUT = ROOT / "docs" / "EMS_Modbus双从站读写模拟测试指导书_V1.7.docx"

IMAGES = {
    "acdc_fault": Path(r"C:\Users\DJ\AppData\Local\Temp\codex-clipboard-f81224df-81e7-4b49-b74c-426d4ea372c9.png"),
    "acdc_recovery": Path(r"C:\Users\DJ\AppData\Local\Temp\codex-clipboard-eb04f9d9-cd9c-4e43-8d98-7eef6d915afd.png"),
    "acdc_stable": Path(r"C:\Users\DJ\AppData\Local\Temp\codex-clipboard-658d360f-e7be-4497-8597-861a0ad1bc7c.png"),
    "dcdc_fault": Path(r"C:\Users\DJ\AppData\Local\Temp\codex-clipboard-b0168d64-03a3-4370-9452-155b3968fa62.png"),
    "dcdc_recovery": Path(r"C:\Users\DJ\AppData\Local\Temp\codex-clipboard-55b2ffb8-28ba-4698-94e1-daabcedfa8df.png"),
}


def set_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_margins(cell, top=100, start=120, bottom=100, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def format_table(table, widths_cm):
    table.autofit = False
    for row_index, row in enumerate(table.rows):
        for col_index, cell in enumerate(row.cells):
            cell.width = Cm(widths_cm[col_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_margins(cell)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.05
                for run in paragraph.runs:
                    run.font.name = "Microsoft YaHei"
                    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "微软雅黑")
                    run.font.size = Pt(9)
            if row_index == 0:
                set_shading(cell, "D9EAF7")
                for run in cell.paragraphs[0].runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor(31, 78, 121)


def add_bullet(doc, text):
    doc.add_paragraph(text, style="List Bullet")


def add_picture(doc, key, width_cm, caption):
    path = IMAGES[key]
    if not path.exists():
        raise FileNotFoundError(path)
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run().add_picture(str(path), width=Cm(width_cm))
    paragraph = doc.add_paragraph(caption)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(8)


def add_result_table(doc, rows):
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for index, value in enumerate(("观察项", "异常期间", "取消异常后", "判定")):
        table.rows[0].cells[index].text = value
    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = value
    format_table(table, [4.0, 4.5, 4.8, 3.1])


doc = Document(SOURCE)
doc.core_properties.title = "EMS Modbus RTU 双从站 PC 模拟测试指导书 V1.7"
doc.add_page_break()
doc.add_heading("21. Modbus异常响应码隔离与自动恢复测试", level=1)

p = doc.add_paragraph()
p.add_run("测试日期：").bold = True
p.add_run("2026年7月28日    ")
p.add_run("异常类型：").bold = True
p.add_run("Exception 06（Slave Device Busy）    ")
p.add_run("记录版本：").bold = True
p.add_run("V1.7")

doc.add_heading("21.1 测试目的与操作方法", level=2)
add_bullet(doc, "验证EMS能够把Modbus异常响应识别为协议错误，而不是无响应超时。")
add_bullet(doc, "验证ACDC或DCDC单独返回异常响应时，另一从站仍能持续轮询且错误计数不增加。")
add_bullet(doc, "验证取消异常响应后，目标从站无需复位EMS即可自动恢复。")
add_bullet(doc, "历史错误计数恢复后不清零；判定依据是当前状态恢复、错误停止增长且成功计数继续增加。")

p = doc.add_paragraph()
p.add_run("配置：").bold = True
p.add_run(
    "ACDC使用ID 33、功能03、地址4096、数量13窗口；DCDC使用ID 1、功能03、地址1028、数量11窗口。"
    "目标窗口仅勾选Return exception 06, Busy，Skip response和Insert CRC/LRC error均不勾选。"
    "异常注入与取消期间EMS和Keil Debug保持运行。"
)

doc.add_heading("21.2 ACDC Exception 06注入、DCDC隔离与恢复", level=2)
add_result_table(doc, [
    ("ACDC online", "0", "1", "自动恢复"),
    ("ACDC consecutive_failures", "0x60 = 96", "0", "当前失败清除"),
    ("ACDC last_error", "0xFFFD = -3", "0", "异常响应识别正确"),
    ("ACDC success_count", "0x1B2 = 434", "0x230 = 560；稳定复核0x274 = 628", "恢复后持续增加"),
    ("ACDC timeout_count", "0", "0", "未误判为超时"),
    ("ACDC protocol_error_count", "0x71 = 113", "过渡至0x75 = 117，随后保持不变", "异常停止后稳定"),
    ("DCDC success_count", "0x21C = 540", "0x29E = 670；稳定复核0x2E2 = 738", "全程持续轮询"),
    ("DCDC错误计数", "timeout=0，protocol=0x07", "保持不变", "ACDC→DCDC隔离通过"),
])

p = doc.add_paragraph()
p.add_run("过渡计数说明：").bold = True
p.add_run(
    "取消Return exception后，ACDC protocol_error_count由0x71短暂增加到0x75，"
    "可能来自已经发出或正在处理的请求。继续运行30秒后仍保持0x75，"
    "同时ACDC与DCDC成功计数各增加68，确认恢复已稳定。"
)

add_picture(doc, "acdc_fault", 11.0, "图29  ACDC返回Exception 06时离线并记录协议错误，DCDC保持正常")
add_picture(doc, "acdc_recovery", 11.0, "图30  取消ACDC异常后的初始恢复状态")
add_picture(doc, "acdc_stable", 11.0, "图31  继续运行30秒后ACDC错误计数稳定，双从站成功计数继续增加")

doc.add_heading("21.3 DCDC Exception 06注入、ACDC隔离与恢复", level=2)
add_result_table(doc, [
    ("DCDC online", "1", "1", "保持在线"),
    ("DCDC consecutive_failures", "0", "0", "有效响应会清除连续失败"),
    ("DCDC last_error", "0", "0", "截图时当前事务正常"),
    ("DCDC success_count", "0x32B = 811", "0x37B = 891", "取消后增加80"),
    ("DCDC timeout_count", "0", "0", "未误判为超时"),
    ("DCDC protocol_error_count", "由0x07增至0x0F = 15", "保持0x0F", "识别8次异常响应"),
    ("ACDC success_count", "0x2C5 = 709", "0x316 = 790", "全程增加81"),
    ("ACDC错误计数", "timeout=0，protocol=0x75", "保持不变", "DCDC→ACDC隔离通过"),
])

p = doc.add_paragraph()
p.add_run("在线状态说明：").bold = True
p.add_run(
    "本轮DCDC异常响应呈间歇注入，正常响应会把consecutive_failures清零，"
    "因此未达到连续3次失败的离线门限。验收重点为protocol_error_count增加、"
    "timeout_count不增加、另一从站不受影响以及取消异常后错误停止增长，上述条件均满足。"
)

add_picture(doc, "dcdc_fault", 11.0, "图32  DCDC Exception 06注入后协议错误增至0x0F，ACDC保持正常")
add_picture(doc, "dcdc_recovery", 11.0, "图33  取消DCDC异常后protocol保持0x0F，双从站成功计数继续增加")

doc.add_heading("21.4 验收结论", level=2)
table = doc.add_table(rows=1, cols=3)
table.style = "Table Grid"
for index, value in enumerate(("验收项", "实测结果", "结论")):
    table.rows[0].cells[index].text = value
for values in [
    ("ACDC异常响应识别", "last_error=-3，protocol增加，timeout为0", "通过"),
    ("ACDC异常时DCDC隔离", "DCDC成功计数持续增加，错误计数不变", "通过"),
    ("ACDC取消异常后恢复", "无需复位，状态恢复且protocol稳定在0x75", "通过"),
    ("DCDC异常响应识别", "protocol由0x07增至0x0F，timeout为0", "通过"),
    ("DCDC异常时ACDC隔离", "ACDC成功计数持续增加，错误计数不变", "通过"),
    ("DCDC取消异常后恢复", "protocol保持0x0F，成功计数继续增加", "通过"),
]:
    cells = table.add_row().cells
    for index, value in enumerate(values):
        cells[index].text = value
format_table(table, [7.0, 6.2, 3.2])

conclusion = doc.add_table(rows=1, cols=1)
conclusion.style = "Table Grid"
cell = conclusion.cell(0, 0)
cell.text = (
    "最终结论：ACDC和DCDC双向Exception 06异常响应码识别、隔离与自动恢复测试通过。"
    "异常响应被归类到protocol_error_count，未错误计入timeout_count；"
    "任一从站返回异常不会阻塞或污染另一从站，取消异常后无需复位EMS即可恢复稳定通信。"
)
set_shading(cell, "E2F0D9")
set_margins(cell, top=170, start=190, bottom=170, end=190)
for run in cell.paragraphs[0].runs:
    run.bold = True
    run.font.name = "Microsoft YaHei"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(55, 86, 35)

doc.add_heading("21.5 后续测试", level=2)
add_bullet(doc, "执行至少5轮异常注入与恢复重复性测试，逐轮记录目标、异常类型、错误增量、恢复状态和另一从站状态。")
add_bullet(doc, "统计取消异常到online=1、last_error=0及成功计数重新增加所需时间。")
add_bullet(doc, "重复测试期间同步监视Client_Sd[0]底层电池数组，确认通信异常不再触发电池页面异常值。")

doc.save(OUTPUT)
print(OUTPUT)
