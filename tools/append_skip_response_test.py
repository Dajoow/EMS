from pathlib import Path
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(r"C:\Users\DJ\Desktop\Hdu_EMS_H750")
SOURCE = ROOT / "docs" / "EMS_Modbus双从站读写模拟测试指导书_V1.1.docx"
OUTPUT = ROOT / "docs" / "EMS_Modbus双从站读写模拟测试指导书_V1.2.docx"
PHOTO = Path(r"C:\Users\DJ\AppData\Local\Temp\codex-clipboard-97df084d-0d29-4b46-8c95-7142b058e66e.jpg")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def format_table(table, widths_cm):
    table.autofit = False
    for row_index, row in enumerate(table.rows):
        for col_index, cell in enumerate(row.cells):
            cell.width = Cm(widths_cm[col_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.05
                for run in paragraph.runs:
                    run.font.name = "Microsoft YaHei"
                    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "微软雅黑")
                    run.font.size = Pt(9)
            if row_index == 0:
                set_cell_shading(cell, "D9EAF7")
                for run in cell.paragraphs[0].runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor(31, 78, 121)


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(text)
    return p


doc = Document(SOURCE)

doc.add_page_break()
doc.add_heading("17. Skip response 超时异常测试记录", level=1)

p = doc.add_paragraph()
r = p.add_run("测试日期：")
r.bold = True
p.add_run("2026年7月21日    ")
r = p.add_run("测试阶段：")
r.bold = True
p.add_run("DCDC 写通信异常与恢复测试    ")
r = p.add_run("记录版本：")
r.bold = True
p.add_run("V1.2")

doc.add_heading("17.1 测试目的与操作", level=2)
add_bullet(doc, "使用PC端Modbus Slave模拟DCDC从站，在目标从站窗口勾选Skip response，使从站收到请求但不发送响应。")
add_bullet(doc, "观察EMS主站写命令超时检测、错误状态、取消Skip response后的恢复情况，以及整机界面是否受到影响。")
add_bullet(doc, "本轮重点记录异常现象，尚未对工程代码做任何修改。")

doc.add_heading("17.2 实际现象", level=2)
table = doc.add_table(rows=1, cols=3)
table.style = "Table Grid"
for i, value in enumerate(("观察项", "实测结果", "判定")):
    table.rows[0].cells[i].text = value
rows = [
    ("从站响应", "勾选Skip response后，PC模拟从站停止返回目标请求的响应帧。", "符合异常注入预期"),
    ("EMS写状态", "超时路径能够被触发；此前观察到last_error=-2，写状态进入失败状态。", "超时检测有效"),
    ("取消异常后恢复", "取消Skip response后重新触发时曾出现继续失败，需要进一步复测恢复路径。", "恢复测试暂未通过"),
    ("电池信息界面", "界面出现明显不合理的数据，包括8.744V、34.816V、819.2%、-245.76℃、102.74℃等。", "严重异常"),
    ("复位表现", "复位后相关运行数据结构会重新清零；本次未完成复位后完整恢复时间和重复次数统计。", "待补充"),
]
for values in rows:
    cells = table.add_row().cells
    for i, value in enumerate(values):
        cells[i].text = value
format_table(table, [3.2, 10.0, 3.0])

if PHOTO.exists():
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(PHOTO), width=Cm(15.8))
    p = doc.add_paragraph("图10  勾选Skip response期间，EMS电池信息界面出现异常数值")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True

doc.add_heading("17.3 数据特征与初步判断", level=2)
p = doc.add_paragraph()
p.add_run("已确认：").bold = True
p.add_run("界面显示的是经过正常格式化的异常数值，不是LCD像素或文字排版损坏。电池页面直接使用Client_Sd[]中的BAT_VOL、BAT_SOC和BAT_TMP数据，当前显示路径没有在线状态和合理范围拦截。")

p = doc.add_paragraph()
p.add_run("数据特征：").bold = True
p.add_run("819.2%、-245.76℃、34.816V等数值可对应0x2000、0xA000、0x8800一类明显的二进制位模式，提示底层数据可能被覆盖、读取到未完成更新的数据，或读取了无效缓存。")

p = doc.add_paragraph()
p.add_run("尚未确认：").bold = True
p.add_run("目前不能仅凭界面现象认定具体根因。Modbus任务栈不足、数组/内存越界、Client_Sd[]并发更新以及界面缺少有效性保护均需分别验证。Skip response本身按设计只应产生通信超时，不应修改电池数据。")

doc.add_heading("17.4 本轮测试结论", level=2)
conclusion = doc.add_table(rows=1, cols=1)
conclusion.style = "Table Grid"
cell = conclusion.cell(0, 0)
cell.text = (
    "结论：异常注入成功，EMS能够进入Modbus超时失败路径；但超时期间电池信息界面出现大范围异常值，且取消Skip response后的恢复行为尚未稳定验证。"
    "因此，本轮“超时检测”子项通过，“异常隔离与自动恢复”子项不通过，整项异常测试状态记为：待定位/暂不通过。"
)
set_cell_shading(cell, "FCE4D6")
set_cell_margins(cell, top=160, start=180, bottom=160, end=180)
for run in cell.paragraphs[0].runs:
    run.bold = True
    run.font.name = "Microsoft YaHei"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(156, 48, 0)

doc.add_heading("17.5 后续复测记录项", level=2)
table = doc.add_table(rows=1, cols=3)
table.style = "Table Grid"
for i, value in enumerate(("复测项", "需要记录的值", "状态")):
    table.rows[0].cells[i].text = value
rows = [
    ("Modbus任务栈余量", "调试串口输出modbus_Process_minStack；异常前、异常中、恢复后各记录一次。", "未执行"),
    ("底层电池数组", "同时观察Client_Sd[0].BAT_VOL[]、BAT_SOC[]、BAT_TMP[]是否与屏幕同步变乱。", "未执行"),
    ("单窗口异常注入", "仅对1026～1027控制窗口勾选Skip response，确认1028和1051读取窗口保持响应。", "待复测"),
    ("自动恢复", "取消Skip response后记录重新触发次数、sequence、state、last_error和恢复耗时。", "待复测"),
    ("重复性", "至少重复5轮：正常→Skip response→取消→重试；记录每轮是否出现界面异常。", "待复测"),
]
for values in rows:
    cells = table.add_row().cells
    for i, value in enumerate(values):
        cells[i].text = value
format_table(table, [4.1, 9.4, 2.7])

doc.add_heading("17.6 安全约束", level=2)
add_bullet(doc, "问题定位完成前，不将本异常测试固件用于真实DCDC启停或充放电控制。")
add_bullet(doc, "异常测试时只连接EMS板卡与USB-RS485模拟器，避免真实ACDC/DCDC同时接入总线。")
add_bullet(doc, "界面出现超范围电压、SOC或温度时，不把显示值作为控制依据；先停止测试并复位板卡。")

doc.save(OUTPUT)
print(OUTPUT)
